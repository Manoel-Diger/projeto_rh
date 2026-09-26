"""
base_mestra.py
Lê a planilha Base Mestra (Colaborador, Função, Horário, Salário) e interpreta os horários.
Horários marcados como PENDENTE NUNCA são inferidos.
"""
import re
import unicodedata

from openpyxl import load_workbook


def normaliza_nome(txt):
    """Remove acentos/espaços extras e padroniza caixa, para cruzar folha x base."""
    txt = unicodedata.normalize("NFKD", str(txt or ""))
    txt = "".join(c for c in txt if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", txt).strip().upper()


def parse_moeda(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.-]", "", str(v))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


_INTERVALO = re.compile(r"(\d{1,2})\s*h\s*(\d{2})?\s*[–\-]\s*(\d{1,2})\s*h\s*(\d{2})?", re.I)


def interpreta_horario(txt, intervalo_min=60):
    """
    '12h–21h'            -> entrada 12.0, saída 21.0, 8,0 h/dia (desconta 1h de intervalo)
    '08h–17h48'          -> 08:00–17:48, 8,8 h/dia
    '08h–12h / 13h–17h'  -> dois blocos, 8 h/dia (intervalo já é o vão entre blocos)
    'PENDENTE ...'       -> definido = False
    """
    base = {"texto": str(txt or "").strip(), "definido": False, "entrada": None, "saida": None,
            "horas_dia": None, "blocos": [], "turno": "Pendente"}
    if not txt or "PEND" in str(txt).upper():
        return base
    blocos = []
    for h1, m1, h2, m2 in _INTERVALO.findall(str(txt)):
        ini = int(h1) + (int(m1) / 60 if m1 else 0)
        fim = int(h2) + (int(m2) / 60 if m2 else 0)
        blocos.append((ini, fim))
    if not blocos:
        return base
    bruto = sum(f - i for i, f in blocos)
    if len(blocos) == 1 and bruto > 6:
        horas = bruto - intervalo_min / 60
    else:
        horas = bruto
    entrada, saida = blocos[0][0], blocos[-1][1]
    if horas <= 5.5:
        turno = "Parcial"
    elif entrada >= 11.5:
        turno = "Tarde"
    else:
        turno = "Manhã/Comercial"
    base.update(definido=True, entrada=entrada, saida=saida, horas_dia=round(horas, 2),
                blocos=[[round(a, 2), round(b, 2)] for a, b in blocos], turno=turno)
    return base


def ler_base_mestra(caminho, abas_aceitas, intervalo_min=60):
    wb = load_workbook(caminho, data_only=True, read_only=True)
    aba = next((n for n in wb.sheetnames if n in abas_aceitas), None) or wb.sheetnames[0]
    ws = wb[aba]
    linhas = list(ws.iter_rows(values_only=True))
    cab = [str(c or "").strip().lower() for c in linhas[0]]

    def col(*nomes):
        for n in nomes:
            for i, c in enumerate(cab):
                if n in c:
                    return i
        raise ValueError(f"Coluna '{nomes[0]}' não encontrada na aba '{aba}'. Cabeçalho lido: {cab}")

    i_nome, i_func = col("colaborador", "nome"), col("função", "funcao", "cargo")
    i_hor, i_sal = col("horário", "horario", "jornada"), col("salário", "salario")
    registros = []
    for r in linhas[1:]:
        if not r or not r[i_nome]:
            continue
        registros.append({
            "nome": str(r[i_nome]).strip(),
            "chave": normaliza_nome(r[i_nome]),
            "funcao": str(r[i_func] or "").strip(),
            "horario": interpreta_horario(r[i_hor], intervalo_min),
            "salario": parse_moeda(r[i_sal]),
        })
    return {"aba": aba, "registros": registros}
